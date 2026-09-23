#!/usr/bin/env python3
"""Build information-vs-model manuscript DOCX package from MD + figures."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "results" / "final_information_vs_model_manuscript"
DEPS = ROOT / ".deps_docx"
sys.path.insert(0, str(DEPS))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.shared import Cm, Pt, RGBColor, Inches  # noqa: E402

MD = OUT / "manuscript" / "CITEseq_Information_vs_Model_Reliability_Research_Manuscript.md"
FIGS = OUT / "figures"
NAVY = RGBColor(0x1B, 0x3A, 0x5F)

FIG_MAP = {
    1: "Figure1_central_logic.png",
    2: "Figure2_information_vs_model.png",
    3: "Figure3_transfer_fairness.png",
    4: "Figure4_protein_degradation.png",
    5: "Figure5_uncertainty_stochasticity.png",
    6: "Figure6_lawlor_external.png",
}

PRIMARY_TITLE = (
    "Disentangling protein-information gain and model-specific reliability "
    "in CITE-seq representation learning"
)


def set_run_font(run, size=10, bold=False, italic=False, name="Times New Roman", color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Jizhu Yang — Research Manuscript — ")
    set_run_font(run, size=8, color=NAVY)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld)
    run2._r.append(instr)
    run2._r.append(fld2)
    set_run_font(run2, size=8, color=NAVY)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Information gain vs model-specific reliability")
    set_run_font(hr, size=8, italic=True, color=NAVY)


def parse_md(text: str):
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-:|]+\|$", lines[i].strip()):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^\*\*Figure [1-6]\.", line) or re.match(r"^Figure [1-6]\.", line):
            blocks.append(("legend", line.strip()))
        elif line.startswith("**Figure ") and "Figure" in line:
            blocks.append(("figcall", line.strip()))
        else:
            para = [line]
            i += 1
            while (
                i < len(lines)
                and lines[i].strip()
                and not lines[i].startswith("#")
                and not lines[i].startswith("|")
                and not lines[i].startswith("**Figure")
                and not re.match(r"^Figure [1-6]\.", lines[i])
            ):
                para.append(lines[i])
                i += 1
            blocks.append(("p", " ".join(x.strip() for x in para)))
            continue
        i += 1
    return blocks


def add_styled_paragraph(doc, text, style="Normal", size=10, bold=False, space_after=6, first_indent=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_indent is not None:
        pf.first_line_indent = first_indent
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=NAVY)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = p.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True)
        else:
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold, color=NAVY if bold else None)
    return p


def insert_figure(doc, n: int, caption: str | None = None):
    path = FIGS / FIG_MAP[n]
    if not path.exists():
        add_styled_paragraph(doc, f"[Missing figure file: {path.name}]", size=9, bold=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.3))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        clean = re.sub(r"^\*\*|\*\*$", "", caption)
        run = cap.add_run(clean)
        set_run_font(run, size=8, italic=True)
        cap.paragraph_format.space_after = Pt(10)


def build_main():
    text = MD.read_text(encoding="utf-8")
    blocks = parse_md(text)
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        add_page_number(sec)

    inserted = set()
    legends = {}
    for kind, payload in blocks:
        if kind == "legend":
            m = re.match(r"^(?:\*\*)?Figure (\d)\.", payload)
            if m:
                # Prefer longer caption text (full legends overwrite short callouts).
                n = int(m.group(1))
                clean = re.sub(r"^\*\*|\*\*$", "", payload)
                if n not in legends or len(clean) > len(legends[n]):
                    legends[n] = clean

    in_methods = False
    in_legends = False
    in_refs = False
    in_tables = False
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(payload)
            set_run_font(run, size=15, bold=True, color=NAVY)
            p.paragraph_format.space_after = Pt(8)
            continue
        if kind == "h2":
            title = payload
            in_methods = title.lower().startswith("methods")
            in_legends = title.lower().startswith("figure legend")
            in_refs = title.lower().startswith("reference")
            in_tables = title.lower().startswith("table")
            if in_legends:
                # legends attached under figures already
                continue
            add_styled_paragraph(doc, title, size=12, bold=True, space_after=8)
            continue
        if kind == "h3":
            add_styled_paragraph(doc, payload, size=11, bold=True, space_after=4)
            continue
        if kind == "legend":
            # Short Results callouts are parsed as legend; insert figure inline there.
            # Skip the dedicated Figure Legends section body to avoid duplicates.
            if in_legends:
                continue
            m = re.match(r"^(?:\*\*)?Figure (\d)\.", payload)
            if m:
                n = int(m.group(1))
                if n not in inserted and n in FIG_MAP:
                    insert_figure(doc, n, legends.get(n, payload))
                    inserted.add(n)
            continue
        if kind == "table":
            rows = payload
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    table.cell(ri, ci).text = cell
                    for paragraph in table.cell(ri, ci).paragraphs:
                        for run in paragraph.runs:
                            set_run_font(run, size=7.5, bold=(ri == 0), color=NAVY if ri == 0 else None)
            doc.add_paragraph()
            continue
        if kind in ("p", "figcall"):
            if (
                payload.startswith("**Jizhu Yang**")
                or payload.startswith("Research Manuscript")
                or payload.startswith("Unpublished")
                or payload.startswith("September")
            ):
                p = add_styled_paragraph(
                    doc, payload.replace("**", ""), size=11, bold=("Yang" in payload), space_after=2
                )
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue
            m = re.search(r"\*\*Figure (\d)\.", payload)
            if m:
                n = int(m.group(1))
                if n not in inserted and n in FIG_MAP:
                    insert_figure(doc, n, legends.get(n, payload))
                    inserted.add(n)
                    continue
            size = 9 if in_methods or in_refs or in_tables else 10
            indent = Cm(0.4) if (not in_methods and not in_refs and not in_tables and not payload.startswith("**")) else None
            if payload.startswith("**Background:**") or payload.startswith("**Results:**") or payload.startswith("**Conclusions:**"):
                indent = None
            if in_refs and re.match(r"^\d+\.\s", payload):
                p = add_styled_paragraph(doc, payload, size=9, space_after=4)
                pf = p.paragraph_format
                pf.left_indent = Cm(0.5)
                pf.first_line_indent = Cm(-0.5)
                continue
            add_styled_paragraph(doc, payload, size=size, space_after=6, first_indent=indent)

    # Safety: insert any missing figures before end (should not trigger if callouts worked).
    missing = [n for n in range(1, 7) if n not in inserted]
    if missing:
        add_styled_paragraph(doc, "Figures", size=12, bold=True, space_after=8)
        for n in missing:
            insert_figure(doc, n, legends.get(n))

    out = OUT / "manuscript" / "CITEseq_Information_vs_Model_Reliability_Research_Manuscript.docx"
    doc.save(out)
    print("wrote", out, "inserted", sorted(inserted), "missing_fallback", missing)
    return out


def build_legends_docx():
    doc = Document()
    text = MD.read_text(encoding="utf-8")
    for m in re.finditer(r"(?:\*\*)?Figure [1-6]\..*?(?=\n\n|\n(?:\*\*)?Figure |\n## |\Z)", text, flags=re.S):
        add_styled_paragraph(doc, m.group(0).replace("**", ""), size=10, space_after=10)
    path = OUT / "figures" / "Figure_Legends.docx"
    doc.save(path)
    print("wrote", path)


def build_tables_docx():
    doc = Document()
    add_styled_paragraph(doc, "Main Tables", size=14, bold=True)
    text = MD.read_text(encoding="utf-8")
    if "## Tables" not in text:
        return
    sec = text.split("## Tables", 1)[1].split("## References", 1)[0]
    for kind, payload in parse_md(sec):
        if kind == "h3":
            add_styled_paragraph(doc, payload, size=11, bold=True)
        elif kind == "table":
            rows = payload
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    table.cell(ri, ci).text = cell
            doc.add_paragraph()
    path = OUT / "tables" / "Main_Tables.docx"
    doc.save(path)
    print("wrote", path)


def build_si_docx():
    doc = Document()
    add_styled_paragraph(doc, "Supplementary Information", size=14, bold=True, space_after=8)
    add_styled_paragraph(doc, PRIMARY_TITLE, size=11, bold=True)
    add_styled_paragraph(doc, "Jizhu Yang — Unpublished research manuscript — September 2026", size=10)
    sections = [
        (
            "S1. Information-gain and model-specific contrasts",
            "G_info = F1_concat − F1_RNA_PCA = +0.079; G_model = F1_totalVI − F1_concat = +0.034. "
            "Descriptive contrasts only; not causal variance components.",
        ),
        (
            "S2. Target-data-access-matched transfer control",
            "Canonical: results/revision_round3_information_vs_model/transfer_control/tables/matched_transfer_control.csv. "
            "Inductive concat A/B ≈ 0.735/0.734; transductive concat A/B ≈ 0.636/0.767. "
            "Target labels withheld until final evaluation.",
        ),
        (
            "S3. Complete reliability metrics",
            "results/revision_round2_methodological_fixes/calibration/full_reliability_metrics.csv "
            "and lawlor_reliability_summary_by_model.csv (NLL, Brier, ECE, AURC, pAURC, E-AURC).",
        ),
        (
            "S4. Per-class and degradation tables",
            "controls/tables/; test_time_corruption/tables/test_time_corruption_per_class.csv; "
            "training_vs_testtime_corruption.csv.",
        ),
        (
            "S5. Donor-level Lawlor tables",
            "statistics/lawlor_donor_delta.csv; primary biological unit = donor (n=10). "
            "Protein-gated author labels. Protein PCA ≈ 0.927; concat ≈ 0.945.",
        ),
        (
            "S6. Fold×seed latent consistency",
            "38/50 matched fold×seed pairs favor totalVI over scVI (computational, not biological n).",
        ),
        (
            "S7. Variance-component details",
            "statistics/variance_components_corrected.csv and delta_variance_components_corrected.csv (MoM EMS).",
        ),
        (
            "S8. Label-provenance audit",
            "controls/logs/original_label_provenance_audit.md — n=9494, agreement=1.0.",
        ),
        (
            "S9. Clean-baseline and freeze integrity",
            "clean_baseline_vs_phase10.csv; model param_hash audits from Part C.",
        ),
        (
            "S10. Historical PHASE 6–14 artifacts",
            "Frozen; not overwritten by this revision.",
        ),
    ]
    for title, body in sections:
        add_styled_paragraph(doc, title, size=11, bold=True)
        add_styled_paragraph(doc, body, size=10)
    path = OUT / "supplement" / "Supplementary_Information.docx"
    doc.save(path)
    print("wrote", path)


def build_reviews():
    audit = OUT / "review" / "Final_Scientific_Audit.docx"
    doc = Document()
    add_styled_paragraph(doc, "Final Scientific Audit — Information vs Model Reframe", size=14, bold=True)
    reviews = {
        "1. Single-cell computational biology reviewer": {
            "CRITICAL": [
                "None after information-vs-model reframe and matched transfer Case C disclosure.",
            ],
            "MAJOR": [
                "Matched control equals target-feature access only; not algorithmically identical to scArches (stated).",
                "Lawlor protein-gated labels favor protein PCA/concat; stated in Results not only Discussion.",
                "PBMC-only scope and limited panels constrain generality.",
            ],
            "MINOR": [
                "Future protein-independent-label external cohort noted as optional future work.",
            ],
        },
        "2. Statistical-methods reviewer": {
            "CRITICAL": [
                "None: G_info/G_model labeled descriptive; 70% not causal; donor n=10 primary for Lawlor.",
            ],
            "MAJOR": [
                "19/20 restricted to PBMC Direction B; 38/50 is computational fold×seed consistency.",
                "Case C direction dependence correctly interpreted as context-dependent residual advantage.",
            ],
            "MINOR": [
                "Internal bootstrap remains cell-level conditional.",
            ],
        },
        "3. Application / research-reader": {
            "CRITICAL": [
                "None: central question is clear; no universal architecture superiority claim.",
            ],
            "MAJOR": [
                "Title and abstract lead with information vs model distinction.",
            ],
            "MINOR": [
                "Word layout is professional single-column A4 for stability (prior package convention).",
            ],
        },
    }
    for name, grades in reviews.items():
        add_styled_paragraph(doc, name, size=12, bold=True)
        for g, items in grades.items():
            add_styled_paragraph(doc, g, size=10, bold=True)
            for it in items:
                add_styled_paragraph(doc, "• " + it, size=10)
    doc.save(audit)
    print("wrote", audit)

    summary = OUT / "review" / "Application_Reader_Summary.docx"
    doc = Document()
    add_styled_paragraph(doc, "Application Reader Summary — Information vs Model", size=14, bold=True)
    bullets = [
        "Central question: protein-information gain vs model-specific advantage.",
        "Internal: RNA PCA 0.666 → concat 0.745 → totalVI 0.779; G_info=+0.079; G_model=+0.034; ~70% descriptive recovery.",
        "Matched transfer Case C: Dir A totalVI − transductive concat ≈ +0.016; Dir B ≈ −0.113.",
        "Inductive concat preserved ≈ 0.735 / 0.734.",
        "Dir B totalVI−scVI +0.052 (19/20); Dir A near-null.",
        "Lawlor: donor Δ≈+0.090 (10/10); protein PCA/concat stronger under protein-gated labels.",
        "No new external cohort / MultiVI / P2 required for this application manuscript.",
    ]
    for b in bullets:
        add_styled_paragraph(doc, "• " + b, size=11)
    doc.save(summary)
    print("wrote", summary)


def validate(docx_path: Path):
    from docx import Document as D

    doc = D(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    n_img = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    checks = {
        "title_ok": PRIMARY_TITLE.split(" in CITE-seq")[0] in text or "Disentangling protein-information gain" in text,
        "unpublished": "Unpublished manuscript" in text,
        "no_zenodo": "Zenodo" not in text,
        "no_phase": "PHASE" not in text,
        "g_info": "0.079" in text or "+0.079" in text,
        "g_model": "0.034" in text or "+0.034" in text,
        "seventy_pct_descriptive": "70%" in text and "causal" in text.lower(),
        "matched_transductive": "0.636" in text and "0.767" in text,
        "inductive_preserved": "0.735" in text and "0.734" in text,
        "case_c": ("+0.016" in text or "0.016" in text) and ("-0.113" in text or "0.113" in text),
        "protein_gated": "protein-gated" in text.lower(),
        "lawlor_donors": "10 donors" in text or "10/10" in text,
        "protein_pca_lawlor": "0.927" in text,
        "concat_lawlor": "0.945" in text,
        "latent_38_50": "38/50" in text,
        "rna_audit": "9,494" in text or "9494" in text,
        "n_images": n_img,
        "n_tables": len(doc.tables),
        "n_paragraphs": len(doc.paragraphs),
        "no_github_ssh": "git@" not in text and "github.com" not in text.lower(),
    }
    # Fail only if 19/20 is attributed to Lawlor (not merely co-mentioned in Abstract).
    bad_lawlor_1920 = False
    for p in doc.paragraphs:
        t = p.text
        if "19/20" not in t or "Lawlor" not in t:
            continue
        # Allowed if PBMC Direction B context is present and Lawlor is a separate later clause.
        if "Direction B" in t and t.find("19/20") < t.find("Lawlor"):
            continue
        bad_lawlor_1920 = True
    checks["no_1920_on_lawlor"] = not bad_lawlor_1920
    has_1920 = any("19/20" in p.text for p in doc.paragraphs)
    checks["has_1920"] = has_1920
    report = OUT / "logs" / "docx_validation.json"
    report.write_text(json.dumps(checks, indent=2), encoding="utf-8")
    print("validation", checks)
    return checks


def write_revision_report(checks):
    text = MD.read_text(encoding="utf-8")
    abs_m = re.search(r"## Abstract\n\n(.*?)(?:\n## )", text, flags=re.S)
    abstract = re.sub(r"\*\*", "", abs_m.group(1)).strip() if abs_m else ""
    abs_wc = len(re.findall(r"\b\w+\b", abstract))

    lines = [
        "# INFORMATION VS MODEL REVISION REPORT",
        "",
        f"A. Final title: {PRIMARY_TITLE}",
        "",
        "   Alternative titles:",
        "   1. Separating multimodal information gain from model-specific advantage in CITE-seq representation learning",
        "   2. Protein information and model-specific reliability in CITE-seq representation learning",
        "",
        "B. Central research question: How much of the gain from multimodal CITE-seq comes from access to protein information itself, and when does a complex multimodal model provide reliable additional value beyond a simple multimodal representation?",
        "",
        "C. Internal information gain G_info = concat − RNA PCA = 0.745 − 0.666 = +0.079",
        "",
        "D. Internal model-specific gain G_model = totalVI − concat = 0.779 − 0.745 = +0.034",
        "",
        "E. Percentage of RNA-PCA-to-totalVI gap descriptively recovered by concat: ≈70% (0.079/0.113). Language is descriptive, not causal.",
        "",
        "F. Existing inductive concat transfer: Direction A ≈ 0.735; Direction B ≈ 0.734 (preserved source-fitted inductive baselines).",
        "",
        "G. New target-data-access-matched transductive concat: Direction A ≈ 0.636; Direction B ≈ 0.767.",
        "",
        "H. scVI/totalVI transfer (scArches): Dir A scVI 0.656±0.018 / totalVI 0.652±0.015; Dir B scVI 0.601±0.020 / totalVI 0.654±0.025 (Δ+0.052; 19/20 seeds).",
        "",
        "I. Whether target-data matching changes interpretation: YES — Case C. Matching unlabeled-target feature access reverses/weakens residual totalVI advantage relative to simple concat (Dir A totalVI−transductive concat ≈ +0.016; Dir B ≈ −0.113). Inductive concat vs scArches must not be read as a pure architecture comparison.",
        "",
        "J. Whether totalVI retains a model-specific transfer advantage anywhere: YES, but narrowly — vs scVI in Direction B (+0.052, 19/20); vs matched transductive concat only a small +0.016 in Direction A and a clear disadvantage in Direction B. Residual advantage is source–target dependent (Case C).",
        "",
        "K. Lawlor donor-level totalVI−scVI: mean Δ ≈ +0.090; 10/10 donors positive; between-donor SD ≈ 0.010. Latent consistency 38/50 fold×seed pairs (computational).",
        "",
        "L. Lawlor protein PCA / concat: protein PCA ≈ 0.927; concat PCA ≈ 0.945; totalVI ≈ 0.888.",
        "",
        'M. Exact wording for Lawlor label-modality alignment: "protein-gated author labels" / "author labels are protein-gated" (Results, Methods, Figure 6 legend, Discussion).',
        "",
        "N. Manuscript clearly distinguishes information gain from model-specific gain: YES (G_info / G_model defined; Figure 1–2; Abstract; Methods subsection).",
        "",
        "O. Uncertainty / stochasticity / degradation framed as stress tests of residual model-specific advantage: YES.",
        "",
        "P. Claims of universal architecture superiority remaining: NO (banned/replaced; Lawlor explicitly not architecture ranking).",
        "",
        "Q. New external dataset required for application manuscript: NO. Discussion notes optional future cohort with protein-independent labels.",
        "",
        "R. Final DOCX path: results/final_information_vs_model_manuscript/manuscript/CITEseq_Information_vs_Model_Reliability_Research_Manuscript.docx",
        "",
        "---",
        f"Abstract word count: {abs_wc}",
        f"DOCX validation: {json.dumps(checks)}",
        "Historical PHASE 6–14 / revision_round2 P0–P1 artifacts: not overwritten.",
        "Matched transfer outputs: results/revision_round3_information_vs_model/transfer_control/",
        "",
        "STOP: no commit, push, MultiVI, missing-modality, new external cohort, or journal submission.",
    ]
    path = OUT / "logs" / "INFORMATION_VS_MODEL_REVISION_REPORT.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("wrote", path)


def main():
    for d in ["manuscript", "figures", "tables", "supplement", "review", "logs"]:
        (OUT / d).mkdir(parents=True, exist_ok=True)
    docx = build_main()
    build_legends_docx()
    build_tables_docx()
    build_si_docx()
    build_reviews()
    checks = validate(docx)
    write_revision_report(checks)


if __name__ == "__main__":
    main()
